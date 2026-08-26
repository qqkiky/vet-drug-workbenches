# -*- coding: utf-8 -*-
"""
兽药注册工作台 - 正式流水线
1) 全量回补农业农村部公告（列表 28 页 -> 详情页）
2) 解析新兽药/进口/变更三类注册目录，归一化字段
3) 按药名把每条划分为 化药 / 生物制品 / 诊断制品
4) 写入 SQLite (drugs.db)
5) 生成 Excel 工作簿：
     - 化药 / 生物制品 / 诊断制品 三个分类 sheet
     - 统计汇总 sheet：按企业、按年度、按月度 分别统计三类数量
带本地缓存(pages/)，可断点重跑；礼貌延时避免被拦。
"""
import os, re, json, sqlite3, time, urllib.request, urllib.parse
from collections import defaultdict, Counter
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = "http://www.ivdc.org.cn/xxgk/zcfg/nyncbgg/"
ROOT = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(ROOT, "pages")
DB = os.path.join(ROOT, "drugs.db")
XLSX = os.path.join(ROOT, "兽药注册工作台.xlsx")
PET_XLSX = os.path.join(ROOT, "宠物药注册工作台.xlsx")
DOM_PET_XLSX = os.path.join(ROOT, "国产宠物药注册工作台.xlsx")
IMP_PET_XLSX = os.path.join(ROOT, "进口宠物药注册工作台.xlsx")
os.makedirs(CACHE, exist_ok=True)

UA = {"User-Agent": "Mozilla/5.0 (workbench-pipeline)"}
DRUG_TITLE_KW = ["新兽药", "注册", "再注册", "变更", "进口"]

# ---------- 分类关键词 ----------
DIAG_KW = ["试剂盒", "试纸", "检测卡", "诊断试剂", "ELISA", "PCR", "胶体金",
           "荧光免疫", "荧光微球", "抗原检测", "抗体检测", "检测试剂", "免疫层析",
           "检测板", "诊断制品"]
BIO_KW = ["疫苗", "血清", "抗血清", "卵黄抗体", "蛋黄抗体", "抗体", "干扰素",
          "类毒素", "菌苗", "免疫治疗", "卵黄", "亚单位", "白蛋白", "融合蛋白"]
PET_NAME_KW = ["犬", "猫", "宠物", "伴侣动物", "伴侣"]
PET_ORG_KW = ["宠物", "至宠", "爱宠", "毛孩子", "伴侣"]

# 第4条·宠物用成分词典（伴侣动物专用药活性成分）。
# 这些药品用通用成分名，药名不含"犬/猫"，公告详情页也不刊登靶动物（已核实 0/424 页含"靶动物"），
# 只能靠活性成分判定。命中即判宠物药，species 据成分推断。
# 仅收录"明确伴侣动物专用药"成分；人宠共用成分见 AMBIG_ACT_KW，不自动纳入。
PET_ACT_KEYWORDS = [
    ("格拉匹", "犬"),        # 格拉匹仑 graiprant（犬用NSAIDs）
    ("氟雷拉纳", "犬、猫"),   # Bravecto
    ("阿福拉纳", "犬"),       # NexGard
    ("沙罗拉纳", "犬"),       # Simparica
    ("赛拉菌素", "犬、猫"),   # Revolution 大宠爱
    ("米尔贝肟", "犬、猫"),   # 米贝肟（心丝虫预防）
    ("匹莫苯丹", "犬"),       # Vetmedin
    ("马罗匹坦", "犬、猫"),   # Cerenia
    ("米氮平", "猫"),         # 猫食欲刺激
    ("替米沙坦", "猫"),       # 猫高血压
    ("头孢维星", "犬、猫"),   # Convenia
    ("非泼罗尼", "犬、猫"),   # Frontline
    ("吡虫啉", "犬、猫"),     # Advantage
    ("贝那普利", "犬、猫"),   # Fortekor
]
# 人宠共用成分：宠物与 livestock 均用，无靶动物无法定归属，留待人工/外部源判定。
AMBIG_ACT_KW = ["伊维菌素", "美洛昔康", "马波沙星", "莫昔克丁", "多拉菌素", "头孢泊肟", "加巴喷丁"]

# 第1条扩展·适应症参考映射（仅伴侣动物专用药活性成分；非公告原文，来自兽医用药常识，
# 用于工作台参考展示。源公告详情页不刊登适应症，全量覆盖需接外部说明书/药典源）。
INDICATION_MAP = [
    ("格拉匹", "犬骨关节炎疼痛与炎症（非甾体抗炎药）"),
    ("氟雷拉纳", "犬、猫跳蚤与蜱虫等体外寄生虫防治"),
    ("阿福拉纳", "犬跳蚤、蜱虫等体外寄生虫防治"),
    ("沙罗拉纳", "犬跳蚤、蜱虫及螨类等体外寄生虫防治"),
    ("赛拉菌素", "犬、猫体内外寄生虫（蛔虫、钩虫、心丝虫、跳蚤、耳螨等）"),
    ("米尔贝肟", "犬、猫心丝虫预防与肠道线虫防治"),
    ("匹莫苯丹", "犬二尖瓣疾病与充血性心力衰竭"),
    ("马罗匹坦", "犬、猫止吐（疾病、化疗或晕动所致恶心）"),
    ("米氮平", "猫食欲刺激与抗抑郁"),
    ("替米沙坦", "猫高血压、慢性肾病相关蛋白尿"),
    ("头孢维星", "犬、猫皮肤及软组织等细菌感染"),
    ("非泼罗尼", "犬、猫跳蚤、蜱虫等体外寄生虫防治"),
    ("吡虫啉", "犬、猫跳蚤等体外寄生虫防治"),
    ("贝那普利", "犬、猫充血性心力衰竭、高血压及蛋白尿（ACE抑制剂）"),
]


def fetch(rel_url, sleep=0.0, force=False):
    """force=True 时忽略本地缓存强制联网重取（列表页必须用，否则新公告永远看不见）。
    详情页内容一经发布不再变动，仍走缓存。"""
    fname = rel_url.replace("/", "_")
    path = os.path.join(CACHE, fname)
    if os.path.exists(path) and not force:
        return open(path, encoding="utf-8", errors="ignore").read()
    if sleep:
        time.sleep(sleep)
    url = urllib.parse.urljoin(BASE, rel_url)
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=25) as r:
        raw = r.read()
    try:
        html = raw.decode("utf-8")
    except UnicodeDecodeError:
        html = raw.decode("gbk", errors="ignore")
    open(path, "w", encoding="utf-8").write(html)
    return html


def parse_list(html):
    soup = BeautifulSoup(html, "html.parser")
    items = []
    for a in soup.find_all("a", href=True):
        title = a.get("title") or a.get_text(strip=True)
        href = a["href"]
        if not re.search(r"20\d{4}/t20\d{6}_\d+\.htm", href):
            continue
        if "公告" not in title:
            continue
        m = re.search(r"第(\d+)号", title)
        ann = m.group(1) if m else None
        dm = re.search(r"(20\d{4})/t(20\d{6})", href)
        date = dm.group(2)[:4] + "-" + dm.group(2)[4:6] + "-" + dm.group(2)[6:] if dm else ""
        items.append({"ann_num": ann, "title": title, "date": date, "rel_url": href})
    return items


def collect_announcements():
    """列表页强制联网刷新（force=True）。列表页是"有没有新公告"的唯一入口，
    一旦走缓存就会永久停留在首次运行的快照，新公告全部漏检。
    联网失败时回退到本地缓存，保证流程不中断，并计入失败数。"""
    pages = ["index.htm"] + ["index_%d.htm" % i for i in range(1, 28)]
    seen, items = set(), []
    n_fail = 0
    for p in pages:
        try:
            html = fetch(p, sleep=0.15, force=True)
        except Exception as e:
            n_fail += 1
            print("LIST FAIL(联网)", p, e)
            cp = os.path.join(CACHE, p.replace("/", "_"))
            if os.path.exists(cp):
                print("   -> 回退本地缓存", p)
                html = open(cp, encoding="utf-8", errors="ignore").read()
            else:
                continue
        for it in parse_list(html):
            if it["rel_url"] in seen:
                continue
            seen.add(it["rel_url"])
            items.append(it)
    if n_fail:
        print("列表页联网失败 %d/%d 页（已回退缓存，本次增量检测可能不完整）" % (n_fail, len(pages)))
    return items


def classify_table(header_cells):
    h = "".join(header_cells)
    if "新兽药名称" in h:
        return "new_drug"
    if "生产厂名称" in h:
        return "import_drug"
    if "变更事项" in h or "兽药标准来源" in h:
        return "change"
    return None


def normalize(header_cells, row_cells, table_type):
    h = [x.replace("（", "(").replace("）", ")") for x in header_cells]
    d = dict(zip(h, row_cells))
    rec = {"table_type": table_type}
    if table_type == "new_drug":
        rec.update({"drug_name": d.get("新兽药名称", ""), "org": d.get("研制单位", ""),
                    "category": d.get("类别", ""), "cert_no": d.get("新兽药注册证书号", ""),
                    "monitor_period": d.get("监测期", ""), "country": "", "valid_period": "",
                    "change_item": "", "remark": "", "std_source": ""})
    elif table_type == "import_drug":
        rec.update({"drug_name": d.get("兽药名称", ""), "org": d.get("生产厂名称", ""),
                    "country": d.get("国别", ""), "cert_no": d.get("进口兽药注册证书号", ""),
                    "valid_period": d.get("有效期限", ""), "remark": d.get("备注", ""),
                    "category": "", "monitor_period": "", "change_item": "", "std_source": ""})
    elif table_type == "change":
        rec.update({"drug_name": d.get("兽药名称", ""), "org": d.get("申请单位", ""),
                    "std_source": d.get("兽药标准来源", ""), "change_item": d.get("变更事项", ""),
                    "monitor_period": d.get("监测期(针对变更事项)", d.get("监测期（针对变更事项）", "")),
                    "category": "", "country": "", "cert_no": "", "valid_period": "", "remark": ""})
    return rec


def classify_drug_type(name):
    n = name or ""
    for kw in DIAG_KW:
        if kw in n:
            return "诊断制品"
    for kw in BIO_KW:
        if kw in n:
            return "生物制品"
    return "化药"


def _strip_noise(text):
    """移除会触发'犬'误判的'狂犬病'（此处'犬'指狂犬病病毒，并非犬只）。"""
    return (text or "").replace("狂犬病", "")


def detect_species(rec):
    """从药名/备注/变更事项/标准来源/单位中抽取适用动物标签，写入 species 字段（仅作标注，不影响 is_pet 判定）。"""
    blob = _strip_noise(" ".join(str(rec.get(k, "")) for k in
                                 ("drug_name", "remark", "change_item", "std_source", "org")))
    SPECIES_MAP = [("犬", ["犬", "狗"]), ("猫", ["猫"]),
                   ("宠物(伴侣)", ["宠物", "伴侣动物", "伴侣"])]
    tags, seen = [], set()
    for label, kws in SPECIES_MAP:
        if any(k in blob for k in kws) and label not in seen:
            seen.add(label); tags.append(label)
    # 第4条·成分词典补充物种（通用成分名不含犬/猫时）
    for kw, sp in PET_ACT_KEYWORDS:
        if kw in rec.get("drug_name", ""):
            for s in sp.split("、"):
                if s not in seen:
                    seen.add(s); tags.append(s)
    return "、".join(tags)


def detect_indication(name):
    """按活性成分给出参考适应症（仅伴侣动物专用药）；非公告原文，来源兽医用药常识。"""
    n = name or ""
    for kw, ind in INDICATION_MAP:
        if kw in n:
            return ind
    return ""


def classify_kind(name):
    """区分诊断试剂 / 疫苗 / 治疗药，用于适应症来源标注与推导。"""
    n = name or ""
    if any(k in n for k in ["试纸条", "检测", "试剂盒", "ELISA", "PCR", "胶体金", "诊断", "抗体检测"]):
        return "diag"
    if any(k in n for k in ["疫苗", "亚单位", "灭活", "活疫苗", "重组", "基因工程", "瘤苗",
                            "抗血清", "免疫球蛋白", "卵黄抗体", "病毒样", "mRNA"]):
        return "vaccine"
    return "thera"


def derive_indication(name, kind):
    """对诊断试剂/疫苗，从药名推导适应症（用途即写在名称中），标注 derived。

    治疗药的真实适应症需联网查说明书（见 fill_indications_online.py），不在本函数范围。"""
    n = name or ""
    if kind == "diag":
        for sfx in ["胶体金检测试纸条", "检测试纸条", "ELISA抗体检测试剂盒", "抗体检测试剂盒",
                    "检测试剂盒", "试纸条", "试剂盒", "ELISA", "抗体检测"]:
            if n.endswith(sfx):
                return "用于检测" + n[:-len(sfx)]
        return "用于检测" + n
    if kind == "vaccine":
        cut = n.find("疫苗")
        body = n[:cut] if cut > 0 else n
        body = re.sub(r"(四联活|三联活|二联活|二价灭活|三价灭活|二价|三联|四联|二联|灭活|活)$", "", body)
        body = body.rstrip("、-")
        return "用于预防" + body
    return ""


def is_pet(rec):
    # 第1条增强：在原有的 name/remark/org 基础上，新增扫描 change_item 与 std_source，
    # 因为很多宠物适用性写在「变更事项」里（如「增加靶动物犬」）。
    name = rec.get("drug_name", "")
    if "伪狂犬病" in name:
        return False
    blob = _strip_noise(" ".join(str(rec.get(k, "")) for k in
                                 ("drug_name", "remark", "change_item", "std_source")))
    for kw in PET_NAME_KW:
        if kw in blob:
            return True
    for kw in PET_ORG_KW:
        if kw in rec.get("org", ""):
            return True
    # 第4条·宠物成分词典：通用成分名（不含犬/猫）的伴侣动物专用药
    for kw, _ in PET_ACT_KEYWORDS:
        if kw in name:
            return True
    return False


# ---------- 审计闸门（第2条）----------
# 化药但药名含这些词时，疑似应为生物制品，需人工复核（不静默改判）。
BIO_SUSPECT_KW = ["蛋白", "抗体", "疫苗", "血清", "干扰素", "多肽", "重组",
                  "抗原", "类毒素", "菌苗", "卵黄", "生物"]
# 已知化学药，排除误判（如鞣酸蛋白属化学收敛药）。
CHEM_EXCLUDE_KW = ["鞣酸"]


def audit_records(records):
    """分类与宠物判定后的审计：返回可疑记录清单，供人工复核，而非静默错判。"""
    PET_KW = ["犬", "猫", "宠物", "伴侣动物", "伴侣"]
    findings = []
    for r in records:
        dn = r.get("drug_name", "")
        ci = r.get("change_item", "") or ""
        rm = r.get("remark", "") or ""
        ss = r.get("std_source", "") or ""
        org = r.get("org", "") or ""
        dt = r.get("drug_type", "")
        # 1) 漏项预警：宠物词出现（name/remark/change/std，已剔除“狂犬病”误判）却仍标记非宠物
        blob_pet = _strip_noise(dn + rm + ci + ss)
        if r.get("is_pet") == 0 and any(k in blob_pet for k in PET_KW):
            findings.append({"type": "漏项预警", "ann_num": r.get("ann_num"),
                             "drug_name": dn, "change_item": ci[:80],
                             "hint": "文本含宠物词但 is_pet=0，疑似漏收，需复核"})
        # 2) 单位待核：宠物词仅出现在单位名（已剔除“狂犬病”误判），谨慎起见不自动纳入
        elif r.get("is_pet") == 0 and any(k in _strip_noise(org) for k in PET_KW):
            findings.append({"type": "单位待核", "ann_num": r.get("ann_num"),
                             "drug_name": dn, "org": org,
                             "hint": "单位名含宠物词但 is_pet=0，建议人工确认是否宠物药"})
        # 3) 分类可疑：化药但药名含生物制品指示词（排除已知化学药）
        if dt == "化药" and not any(e in dn for e in CHEM_EXCLUDE_KW) \
           and any(k in dn for k in BIO_SUSPECT_KW):
            findings.append({"type": "分类可疑", "ann_num": r.get("ann_num"),
                             "drug_name": dn, "drug_type": dt,
                             "hint": "化药但药名含生物制品指示词，需人工确认是否应为生物制品"})
        # 4) 人宠共用待核：药名含人宠共用成分但 is_pet=0，需靶动物判定
        if r.get("is_pet") == 0 and any(k in dn for k in AMBIG_ACT_KW):
            findings.append({"type": "人宠共用待核", "ann_num": r.get("ann_num"),
                             "drug_name": dn,
                             "hint": "药名含人宠共用成分，无靶动物无法定归属，建议人工确认"})
    return findings


def write_audit_report(records, findings, path):
    """输出审计报表（Markdown），供人工复核。"""
    from collections import Counter
    by_type = Counter(r["drug_type"] for r in records)
    pet_n = sum(1 for r in records if r["is_pet"])
    fmis = [f for f in findings if f["type"] == "漏项预警"]
    fdep = [f for f in findings if f["type"] == "单位待核"]
    fcls = [f for f in findings if f["type"] == "分类可疑"]
    L = []
    L.append("# 兽药注册工作台 · 审计报表")
    L.append("")
    L.append("## 一、总体统计")
    L.append("- 总条目数：%d" % len(records))
    L.append("- 化药：%d　生物制品：%d　诊断制品：%d" % (
        by_type.get("化药", 0), by_type.get("生物制品", 0), by_type.get("诊断制品", 0)))
    L.append("- 宠物相关：%d" % pet_n)
    L.append("")
    L.append("## 二、漏项预警（%d 条）" % len(fmis))
    if fmis:
        for f in fmis:
            L.append("- [%s] %s ｜ %s" % (f["ann_num"], f["drug_name"], f["change_item"]))
    else:
        L.append("- 无（宠物判定增强后未再发现漏项，闸门校验通过）")
    L.append("")
    L.append("## 三、单位待核（%d 条，未自动纳入，建议人工确认）" % len(fdep))
    if fdep:
        for f in fdep:
            L.append("- [%s] %s ｜ 单位：%s" % (f["ann_num"], f["drug_name"], f["org"]))
    else:
        L.append("- 无")
    L.append("")
    L.append("## 四、分类可疑（%d 条，需人工确认是否应为生物制品）" % len(fcls))
    if fcls:
        for f in fcls:
            L.append("- [%s] %s ｜ 当前：%s" % (f["ann_num"], f["drug_name"], f["drug_type"]))
    else:
        L.append("- 无")
    L.append("")
    famb = [f for f in findings if f["type"] == "人宠共用待核"]
    L.append("## 五、人宠共用待核（%d 条，无靶动物无法定归属，建议人工确认）" % len(famb))
    if famb:
        for f in famb:
            L.append("- [%s] %s" % (f["ann_num"], f["drug_name"]))
    else:
        L.append("- 无")
    L.append("")
    open(path, "w", encoding="utf-8").write("\n".join(L))
    return len(fmis), len(fdep), len(fcls), len(famb)


def parse_details(items):
    """返回 (records, done_items)。done_items 为 [(item, 本公告解析出的条目数)]，
    只含抓取成功的公告；抓取失败的不登记，下次运行会重试（每次运行最多试一次）。"""
    records = []
    done_items = []
    n_fail = 0
    n = 0
    for it in items:
        n += 1
        try:
            html = fetch(it["rel_url"], sleep=0.2)
        except Exception as e:
            n_fail += 1
            print("DETAIL FAIL", it["rel_url"], e)
            continue
        before = len(records)
        soup = BeautifulSoup(html, "html.parser")
        # 从详情页补全更完整的公告标题（列表里的老公告标题是简写）
        full = soup.find(string=re.compile(r"^农业农村部公告.*（"))
        ann_title = full.strip() if full else it["title"]
        for t in soup.find_all("table"):
            rows = t.find_all("tr")
            if not rows:
                continue
            header = [c.get_text(strip=True) for c in rows[0].find_all(["td", "th"])]
            ttype = classify_table(header)
            if not ttype:
                continue
            for r in rows[1:]:
                cells = [c.get_text(strip=True) for c in r.find_all(["td", "th"])]
                if not any(cells):
                    continue
                cells = (cells + [""] * len(header))[:len(header)]
                rec = normalize(header, cells, ttype)
                rec["ann_num"] = it["ann_num"]
                rec["ann_date"] = it["date"]
                rec["ann_title"] = ann_title
                rec["source_url"] = urllib.parse.urljoin(BASE, it["rel_url"])
                rec["drug_type"] = classify_drug_type(rec["drug_name"])
                rec["is_pet"] = is_pet(rec)
                rec["species"] = detect_species(rec)
                rec["indication"] = detect_indication(rec["drug_name"])
                rec["indication_src"] = "ref" if rec["indication"] else ""
                # 进口兽药：生产厂名称列缺失/被错填为证书号时，归并为未标注，并尝试回填证书号
                if rec["table_type"] == "import_drug":
                    if rec["org"] and "外兽药证字" in rec["org"]:
                        if not rec["cert_no"]:
                            rec["cert_no"] = rec["org"]
                        rec["org"] = "（进口·生产厂未标注）"
                    elif not rec["org"]:
                        rec["org"] = "（进口·生产厂未标注）"
                records.append(rec)
        done_items.append((it, len(records) - before))
        if n % 25 == 0:
            print("  已处理 %d 条公告, 当前条目 %d" % (n, len(records)))
    if n_fail:
        print("详情页抓取失败 %d 个（未登记，下次运行自动重试）" % n_fail)
    return records, done_items


def save_sqlite(records):
    """增量 UPSERT（第3条）：按自然键 (ann_num, drug_name, table_type, org) 合并，
    已存在的行更新可变字段、不 wiping；库中缺失的行插入。保留历史数据，避免全量重建。"""
    con = sqlite3.connect(DB)
    con.execute("""CREATE TABLE IF NOT EXISTS drugs(
        id INTEGER PRIMARY KEY, ann_num TEXT, ann_date TEXT, ann_title TEXT,
        table_type TEXT, drug_type TEXT, drug_name TEXT, org TEXT, category TEXT,
        country TEXT, cert_no TEXT, monitor_period TEXT, valid_period TEXT,
        change_item TEXT, std_source TEXT, remark TEXT, is_pet INTEGER, source_url TEXT,
        species TEXT, indication TEXT, indication_src TEXT)""")
    cols = [r[1] for r in con.execute("PRAGMA table_info(drugs)")]
    if "species" not in cols:
        con.execute("ALTER TABLE drugs ADD COLUMN species TEXT")
    if "indication" not in cols:
        con.execute("ALTER TABLE drugs ADD COLUMN indication TEXT")
    if "indication_src" not in cols:
        con.execute("ALTER TABLE drugs ADD COLUMN indication_src TEXT")
    for r in records:
        pet = 1 if r["is_pet"] else 0
        sp = r.get("species", "") or ""
        ind = r.get("indication", "") or ""
        isrc = r.get("indication_src", "") or ""
        key = (r["ann_num"], r["drug_name"], r["table_type"], r["org"])
        exist = con.execute(
            "SELECT id FROM drugs WHERE ann_num=? AND drug_name=? AND table_type=? AND org=?",
            key).fetchone()
        if exist:
            con.execute("""UPDATE drugs SET ann_date=?, ann_title=?, drug_type=?, category=?,
                country=?, cert_no=?, monitor_period=?, valid_period=?, change_item=?,
                std_source=?, remark=?, source_url=?, is_pet=?, species=?, indication=?, indication_src=?
                WHERE ann_num=? AND drug_name=? AND table_type=? AND org=?""",
                (r["ann_date"], r["ann_title"], r["drug_type"], r["category"], r["country"],
                 r["cert_no"], r["monitor_period"], r["valid_period"], r["change_item"],
                 r["std_source"], r["remark"], r["source_url"], pet, sp, ind, isrc,
                 r["ann_num"], r["drug_name"], r["table_type"], r["org"]))
        else:
            con.execute("""INSERT INTO drugs
                (ann_num, ann_date, ann_title, table_type, drug_type, drug_name, org,
                 category, country, cert_no, monitor_period, valid_period, change_item,
                 std_source, remark, is_pet, source_url, species, indication, indication_src)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (r["ann_num"], r["ann_date"], r["ann_title"], r["table_type"], r["drug_type"],
                 r["drug_name"], r["org"], r["category"], r["country"], r["cert_no"],
                 r["monitor_period"], r["valid_period"], r["change_item"], r["std_source"],
                 r["remark"], pet, r["source_url"], sp, ind, isrc))
    con.commit()
    con.close()


def load_sqlite():
    """按公告日期倒序读取，保证最新公告永远排在 Excel 顶部。

    历史 bug：此处曾是裸 SELECT（按 rowid/插入顺序返回）。首次全量构建时公告恰好
    按日期倒序入库，导出看起来"像"排好序了；但增量新增的公告 rowid 最大，会被追加
    到表格最末尾——2026-07-30 的新条目因此排在 2018 年数据之后，用户误以为没入表。
    排序键：有日期的在前 → 日期倒序 → 同日按公告号数值倒序 → 同公告内保持原始行序。
    """
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    rows = [dict(r) for r in con.execute("""
        SELECT * FROM drugs
        ORDER BY (ann_date IS NULL OR ann_date = '') ASC,
                 ann_date DESC,
                 CAST(NULLIF(ann_num, '') AS INTEGER) DESC,
                 id ASC
    """)]
    con.close()
    return rows


def existing_ann_nums():
    """[已弃用于增量判定] 仅保留供诊断脚本使用。
    从 drugs 表反推"已处理公告"是错误口径：不含注册目录表格的公告解析后 0 条目，
    永远不会出现在 drugs 表，于是每次运行都被重新判为"新公告"重复解析。"""
    con = sqlite3.connect(DB)
    s = set(r[0] for r in con.execute(
        "SELECT DISTINCT ann_num FROM drugs WHERE ann_num IS NOT NULL AND ann_num != ''"))
    con.close()
    return s


def _ensure_processed_table(con):
    con.execute("""CREATE TABLE IF NOT EXISTS processed_ann(
        rel_url TEXT PRIMARY KEY, ann_num TEXT, ann_date TEXT,
        ann_title TEXT, n_records INTEGER, parsed_at TEXT)""")


def processed_urls():
    """已成功解析过的公告 rel_url 集合。用 rel_url 而非 ann_num 作键：
    rel_url 含年月路径天然唯一，不受公告号跨年重号影响。"""
    con = sqlite3.connect(DB)
    _ensure_processed_table(con)
    s = set(r[0] for r in con.execute("SELECT rel_url FROM processed_ann"))
    con.commit(); con.close()
    return s


def mark_processed(done_items):
    """登记已成功解析的公告（含解析出 0 条目的），使其不再被重复抓取。
    抓取失败的公告不会传入这里，下次运行会自动重试。"""
    con = sqlite3.connect(DB)
    _ensure_processed_table(con)
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    for it, cnt in done_items:
        con.execute("""INSERT INTO processed_ann
            (rel_url, ann_num, ann_date, ann_title, n_records, parsed_at)
            VALUES (?,?,?,?,?,?)
            ON CONFLICT(rel_url) DO UPDATE SET
              n_records=excluded.n_records, parsed_at=excluded.parsed_at""",
            (it["rel_url"], it["ann_num"], it["date"], it["title"], cnt, ts))
    con.commit(); con.close()


def backfill_indication():
    """幂等回填适应症：仅对 indication 为空的行，按 INDICATION_MAP 成分关键词补参考适应症。"""
    con = sqlite3.connect(DB)
    cols = [r[1] for r in con.execute("PRAGMA table_info(drugs)")]
    if "indication_src" not in cols:
        con.execute("ALTER TABLE drugs ADD COLUMN indication_src TEXT")
    n = 0
    for kw, ind in INDICATION_MAP:
        cur = con.execute(
            "UPDATE drugs SET indication=?, indication_src='ref' WHERE (indication IS NULL OR indication='') AND drug_name LIKE ?",
            (ind, "%" + kw + "%"))
        n += cur.rowcount
    con.commit(); con.close()
    return n


def split_orgs(org):
    if not org:
        return []
    return [o.strip() for o in re.split(r"[、，]", org) if o.strip()]


def build_excel(records, xlsx_path, title, pet_only=False):
    wb = Workbook()
    # 分类 sheet 列定义（pet_only 时去掉冗余的"宠物相关"列）
    base10 = [("公告号", "ann_num"), ("公告日期", "ann_date"), ("药品类型", "drug_type"),
              ("药名", "drug_name"), ("研制/生产单位", "org"), ("类别/国别", None),
              ("证书号", "cert_no"), ("监测期/有效期", None), ("变更事项", "change_item"),
              ("备注", "remark"), ("适应症", "indication"), ("适应症来源", "indication_src")]
    if pet_only:
        cols = base10 + [("适用动物", "species"), ("来源链接", "source_url")]
    else:
        cols = base10 + [("宠物相关", None), ("适用动物", "species"), ("来源链接", "source_url")]
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")

    def fill_cat_or_country(r):
        return r.get("category") or r.get("country") or ""
    def fill_period(r):
        return r.get("monitor_period") or r.get("valid_period") or ""
    def fill_pet(r):
        return "是" if r["is_pet"] else ""

    for type_name in ["化药", "生物制品", "诊断制品"]:
        ws = wb.create_sheet(title=type_name)
        sub = [r for r in records if r["drug_type"] == type_name]
        for c, (label, _) in enumerate(cols, 1):
            cell = ws.cell(1, c, label)
            cell.font = hdr_font; cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")
        for i, r in enumerate(sub, 2):
            vals = []
            for label, key in cols:
                if key is None:
                    if label == "类别/国别":
                        vals.append(fill_cat_or_country(r))
                    elif label == "监测期/有效期":
                        vals.append(fill_period(r))
                    elif label == "宠物相关":
                        vals.append(fill_pet(r))
                    elif label == "适用动物":
                        vals.append(r.get("species", "") or "")
                    else:
                        vals.append("")
                else:
                    vals.append(r.get(key, ""))
            for c, v in enumerate(vals, 1):
                ws.cell(i, c, v)
        ws.freeze_panes = "A2"
        if pet_only:
            widths = [9, 12, 10, 40, 34, 12, 18, 16, 30, 14, 34, 12, 10, 40]
        else:
            widths = [9, 12, 10, 40, 34, 12, 18, 16, 30, 14, 34, 12, 9, 10, 40]
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + c) if c <= 26 else "A"].width = w

    # 统计汇总 sheet
    ws = wb.create_sheet(title="统计汇总")
    bold = Font(bold=True)
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    total = len(records)
    by_type = Counter(r["drug_type"] for r in records)
    ws.cell(2, 1, "总条目数").font = bold
    ws.cell(2, 2, total)
    ws.cell(3, 1, "化药").font = bold; ws.cell(3, 2, by_type.get("化药", 0))
    ws.cell(4, 1, "生物制品").font = bold; ws.cell(4, 2, by_type.get("生物制品", 0))
    ws.cell(5, 1, "诊断制品").font = bold; ws.cell(5, 2, by_type.get("诊断制品", 0))
    row0 = 6
    if not pet_only:
        pet_n = sum(1 for r in records if r["is_pet"])
        ws.cell(row0, 1, "其中宠物相关").font = bold; ws.cell(row0, 2, pet_n)
        row0 += 1
    years = sorted({r["ann_date"][:4] for r in records if r["ann_date"]})
    ws.cell(row0, 1, "覆盖年度").font = bold
    ws.cell(row0, 2, "%s ~ %s" % (years[0], years[-1]) if years else "-")
    row0 += 1

    # 1) 按企业
    comp = defaultdict(lambda: {"化药": 0, "生物制品": 0, "诊断制品": 0, "合计": 0})
    for r in records:
        for o in split_orgs(r["org"]):
            comp[o][r["drug_type"]] += 1
            comp[o]["合计"] += 1
    row = row0 + 1
    ws.cell(row, 1, "一、按企业统计（研制/生产单位，按合计降序）").font = Font(bold=True, size=12)
    row += 1
    for c, h in enumerate(["企业/单位", "化药", "生物制品", "诊断制品", "合计"], 1):
        cell = ws.cell(row, c, h); cell.font = hdr_font; cell.fill = hdr_fill
    comp_sorted = sorted(comp.items(), key=lambda kv: kv[1]["合计"], reverse=True)
    for o, d in comp_sorted:
        row += 1
        ws.cell(row, 1, o); ws.cell(row, 2, d["化药"])
        ws.cell(row, 3, d["生物制品"]); ws.cell(row, 4, d["诊断制品"]); ws.cell(row, 5, d["合计"])

    # 2) 按年度
    row += 2
    ws.cell(row, 1, "二、按年度统计").font = Font(bold=True, size=12)
    row += 1
    for c, h in enumerate(["年度", "化药", "生物制品", "诊断制品", "合计"], 1):
        cell = ws.cell(row, c, h); cell.font = hdr_font; cell.fill = hdr_fill
    ystat = defaultdict(lambda: {"化药": 0, "生物制品": 0, "诊断制品": 0, "合计": 0})
    for r in records:
        y = r["ann_date"][:4] if r["ann_date"] else "未知"
        ystat[y][r["drug_type"]] += 1; ystat[y]["合计"] += 1
    for y in sorted(ystat):
        row += 1
        d = ystat[y]
        ws.cell(row, 1, y); ws.cell(row, 2, d["化药"])
        ws.cell(row, 3, d["生物制品"]); ws.cell(row, 4, d["诊断制品"]); ws.cell(row, 5, d["合计"])

    # 3) 按月度
    row += 2
    ws.cell(row, 1, "三、按月度统计").font = Font(bold=True, size=12)
    row += 1
    for c, h in enumerate(["月度", "化药", "生物制品", "诊断制品", "合计"], 1):
        cell = ws.cell(row, c, h); cell.font = hdr_font; cell.fill = hdr_fill
    mstat = defaultdict(lambda: {"化药": 0, "生物制品": 0, "诊断制品": 0, "合计": 0})
    for r in records:
        m = r["ann_date"][:7] if r["ann_date"] else "未知"
        mstat[m][r["drug_type"]] += 1; mstat[m]["合计"] += 1
    for m in sorted(mstat):
        row += 1
        d = mstat[m]
        ws.cell(row, 1, m); ws.cell(row, 2, d["化药"])
        ws.cell(row, 3, d["生物制品"]); ws.cell(row, 4, d["诊断制品"]); ws.cell(row, 5, d["合计"])

    for c, w in enumerate([42, 10, 12, 12, 10], 1):
        ws.column_dimensions[chr(64 + c)].width = w

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(xlsx_path)


def main():
    print("== 收集公告列表(28页) ==")
    items = collect_announcements()
    print("公告链接总数:", len(items))
    done = processed_urls()
    force_full = os.environ.get("FORCE_FULL") == "1"
    if done and not force_full:
        new_items = [it for it in items if it["rel_url"] not in done]
        print("增量模式：已解析 %d 个公告，本次待解析 %d 个" % (len(done), len(new_items)))
    else:
        new_items = items
        print("全量模式：解析 %d 个公告" % len(items))
    print("== 解析详情页并分类 ==")
    records, done_items = parse_details(new_items)
    print("本次解析条目:", len(records))
    # 登记已解析公告（含 0 条目的），避免下次重复抓取
    if done_items:
        mark_processed(done_items)
        n_empty = sum(1 for _, c in done_items if c == 0)
        print("已登记 %d 个公告为已解析（其中 %d 个不含注册目录，0 条目）" % (len(done_items), n_empty))
    if records:
        # 审计闸门（第2条）：分类/判定后校验，输出可疑记录清单
        try:
            findings = audit_records(records)
            n_mis, n_dep, n_cls, n_famb = write_audit_report(records, findings,
                                                     os.path.join(ROOT, "audit_report.md"))
            print("审计闸门(本次新增)：漏项预警 %d / 单位待核 %d / 分类可疑 %d / 人宠共用待核 %d，详见 audit_report.md"
                  % (n_mis, n_dep, n_cls, n_famb))
        except Exception as e:
            print("审计报表生成失败(不影响主流程):", e)
        print("== 增量 UPSERT 写入 SQLite ==")
        save_sqlite(records)
    else:
        print("无新增条目，跳过写入（FORCE_FULL=1 可强制全量重建）")
    # 回填适应症（幂等，仅对空值按成分词典补参考适应症）
    n_fill = backfill_indication()
    print("适应症回填: %d 行" % n_fill)
    # 从 DB 全量生成 Excel（保证工作台完整、含历史 + 新增）
    print("== 从 SQLite 生成 Excel ==")
    all_records = load_sqlite()
    print("库中总条目:", len(all_records))
    build_excel(all_records, XLSX, "兽药注册工作台 · 统计汇总", pet_only=False)
    pet_records = [r for r in all_records if r["is_pet"]]
    print("宠物相关条目:", len(pet_records))
    build_excel(pet_records, PET_XLSX, "宠物药注册工作台 · 统计汇总", pet_only=True)
    # 按药品来源属性拆分：进口 = import_drug；国产 = new_drug + change
    domestic = [r for r in pet_records if r["table_type"] in ("new_drug", "change")]
    imported = [r for r in pet_records if r["table_type"] == "import_drug"]
    print("  国产宠物药:", len(domestic), " 进口宠物药:", len(imported))
    build_excel(domestic, DOM_PET_XLSX, "国产宠物药注册工作台 · 统计汇总", pet_only=True)
    build_excel(imported, IMP_PET_XLSX, "进口宠物药注册工作台 · 统计汇总", pet_only=True)
    print("完成 ->", XLSX)
    print("完成 ->", PET_XLSX)
    print("完成 ->", DOM_PET_XLSX)
    print("完成 ->", IMP_PET_XLSX)


if __name__ == "__main__":
    main()
